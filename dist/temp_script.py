import json
import openpyxl
import os
import shutil

# write json data to file
def write_data_to_json_file(folder, name, result):
  path = os.path.join(folder, name + '.json')
  with open(path, 'w', encoding='utf-8') as file:
    json.dump(result, file, ensure_ascii=False, indent=2)
  print(f'File "{name}.json" successfully saved!')

# parse excel table to json format
def parse_table_to_json(table_name, table_data, json_locale_data):
  content = {}
  column_headers = table_data[0]
  column_names, column_indices = sort_column_indices(column_headers)

  for row in table_data[1:]:
    row_data = {column_names[i]: row[i] for i in column_indices['data'] if row[i] is not None}
    json_data = {column_names[i]: json.loads(row[i]) for i in column_indices['json'] if row[i] is not None}
    ref_data = {column_names[i]: int(row[i].split(':')[0]) for i in column_indices['refs'] if row[i] is not None}
    content[row[0]] = {**row_data, **json_data, **ref_data}
    for i in column_indices['locales']:
      json_locale_data[row[i-3]] = {'text': row[i], 'column': column_names[i], 'table': table_name}

  return content

# sort column indexes by their type
def sort_column_indices(headers):

  column_indices = {
    'data': [],
    'locales': [],
    'json': [],
    'refs': []
  }
  column_names = []

  for i, header in enumerate(headers):
    column_name, column_type = header.split(':')[0], header.split(':')[1].lower()
    column_names.append(column_name)

    if 'support' in column_type: continue
    elif 'ltext' in column_type: column_indices['locales'].append(i)
    elif 'json' in column_type: column_indices['json'].append(i)
    elif 'ref' in column_type: column_indices['refs'].append(i)
    else: column_indices['data'].append(i)

  return column_names, column_indices

# parse workbook to JSON folder
def parse_wb(wb_path):

  json_content_data = []
  json_locale_data = {}

  wb_handler = openpyxl.load_workbook(wb_path, keep_links=True, keep_vba=True)
  for wb_sheet_name in wb_handler.sheetnames:
    if wb_sheet_name.startswith('@'):
      print(f'Sheet {wb_sheet_name} skipped (starts with the @ symbol)')
      continue

    wb_sheet_handler = wb_handler[wb_sheet_name]
    for table_name, table in wb_sheet_handler.tables.items():
      json_content_data.append(table_name + '.json')
      table_range = wb_sheet_handler[table]
      table_data = [[cell.value for cell in row] for row in table_range]
      table_data_json = parse_table_to_json(table_name, table_data, json_locale_data)
      write_data_to_json_file(wb_result_path_dir, table_name, table_data_json)

  write_data_to_json_file(wb_result_path_dir, '@content.json', json_content_data)
  write_data_to_json_file(wb_result_path_dir, '@locale.json',  json_locale_data)

# run compile script
wb_file_path = r'W:\godot\cassiopeia\cas-data-designer\dist\data-designer.xlsm'
wb_file_path_dir = os.path.dirname(wb_file_path)
wb_file_path_name = os.path.basename(wb_file_path).split('.')[0]
wb_result_path_dir = f'{wb_file_path_dir}\\{wb_file_path_name}'

if os.path.exists(wb_result_path_dir): shutil.rmtree(wb_result_path_dir)
os.mkdir(wb_result_path_dir)

if wb_file_path.endswith('.xlsm') or wb_file_path.endswith('.xlsx') or wb_file_path.endswith('.xls'):
  print(f'File detected {wb_file_path}')
  parse_wb(wb_file_path)

print('All files have been processed.')
input('Press Enter to exit...')
