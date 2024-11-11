import json
import openpyxl
import os

# function for converting excel sheet to JSON
def convert_table_to_json(table_name, table_data, json_locale_data):

  # parsing table headers
  column_names = []
  column_indices_with_data = []
  column_indices_with_locales = []
  column_indices_with_refs = []

  for i, header in enumerate(table_data[0]):

      column_names.append(header.split(":")[0])
      column_type = header.split(":")[1]

      if 'formula' in column_type.lower(): continue 
      elif 'ltext' in column_type.lower(): column_indices_with_locales.append(i)
      elif 'ref' in column_type.lower():   column_indices_with_refs.append(i)
      else: column_indices_with_data.append(i)

  # prepare data
  json_table_data = {}
  for row in table_data[1:]:
    
    # handle sheet data 
    json_table_data[row[0]] = (
      {column_names[i]: row[i] for i in column_indices_with_data} | 
      {column_names[i]: int(row[i].split(":")[0]) for i in column_indices_with_refs}
    )

    # handle locale data 
    for i in column_indices_with_locales:
      json_locale_data[row[i-1]] = {"text":row[i], "column": column_names[i], "table": table_name}

  # write target json
  json_path = os.path.join(output_dir, f'{table_name}.json')
  with open(json_path, 'w', encoding='utf-8') as json_file:
    json.dump(json_table_data, json_file, ensure_ascii=False, indent=2)

  print(f'Table {table_name} successfully transformed into JSON')

# parse workbook excel

def parse_wb(wb_path):
  json_locale_data = {}
  wb_handler = openpyxl.load_workbook(wb_path, keep_links=True, keep_vba=True)

  # parse sheet excel
  for wb_sheet_name in wb_handler.sheetnames:
    
    # ignore sheets starting with the symbol '@'
    if wb_sheet_name.startswith('@'):
      print(f'Sheet {wb_sheet_name} skipped (starts with the @ symbol)')
      continue

    wb_sheet_handler = wb_handler[wb_sheet_name]
    for table_name, table in wb_sheet_handler.tables.items():
      table_range = wb_sheet_handler[table]
      table_data = [[cell.value for cell in row] for row in table_range]
      convert_table_to_json(table_name, table_data, json_locale_data)

  json_locale_path = os.path.join(output_dir, f'locale.json')
  with open(json_locale_path, 'w', encoding='utf-8') as json_file:
    json.dump(json_locale_data, json_file, ensure_ascii=False, indent=2)

# setup workflow
input_dir = os.getcwd()
output_dir = os.getcwd()

# scan directory
for root, dirs, files in os.walk(input_dir):
  for file in files:
    wb_file_path = os.path.join(root, file)

    # skipping temporary files (files starting with '~' or '.')
    if file.startswith('~') or file.startswith('.'): continue

    # skipping files with the wrong extension
    if file.endswith('.xlsx') or file.endswith('.xls'):
      print(f'File detected {wb_file_path}')
      parse_wb(wb_file_path)

print("All files have been processed.")
input("Press Enter to exit...")